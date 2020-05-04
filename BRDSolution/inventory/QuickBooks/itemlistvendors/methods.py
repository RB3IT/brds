## Builtin
import csv
import collections
import pathlib
## Third Party
import openpyxl
## This Module
from BRDSolution.inventory.QuickBooks.itemlistvendors import classes
from BRDSolution.inventory import sql


def generatequickbooksexcelconversion(databasefile,quickbooksexcelfile):
    database = sql.InventoryDatabase(databasefile)
    itemstable = database.tables['items']
    coststable = database.tables['costs']
    databaseitems = itemstable.filter(include={"=":True})
    
    book = openpyxl.load_workbook(quickbooksexcelfile,read_only=True)
    quickbookitems = classes.QuickbooksInventory(book['Sheet1']).items

    
    outputitems = []
    for item in databaseitems:
        itemid = item['itemid']
        description = item['description']

        costentry = coststable.filter(itemid={"=":itemid})[0]
        databasecost = costentry['cost']

        if itemlookup in quickbookitems:
            qbitem = quickbookitems[itemlookup]
            qbitemname = qbitem['Item']
            qbcost = qbitem['Cost']
            qbvendor = qbitem['Preferred Vendor']
        else:
            qbcost = ""
            qbvendor = ""

        outputitems.append(collections.OrderedDict([('ItemId',itemid),('Description',description),
                            ('Original Cost',databasecost),
                            ('Quickbooks Item Name',qbitemname),('Quickboooks Cost',qbcost),('Vendor',qbvendor)]))
    
    return outputitems

def exportconversionresults(results,file):
    headers = results[0].keys()
    with open(file,'w',newline="") as f:
        writer = csv.DictWriter(f,fieldnames=headers)
        writer.writeheader()
        writer.writerows(results)

    return True